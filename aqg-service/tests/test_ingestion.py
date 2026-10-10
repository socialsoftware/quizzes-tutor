import time

import pytest
from fastapi.testclient import TestClient

from app.api.main import create_app
from app.chunking.splitter import split_markdown
from app.config import Settings
from app.db import MaterialRow
from app.ingestion.adapters import (
    FallbackAdapter,
    MarkdownAdapter,
    MarkItDownAdapter,
    ParsedDocument,
    Parsers,
    PymupdfAdapter,
    UnsupportedFormat,
    adapter_for,
)
from app.retrieval.retriever import LexicalRetriever
from app.schema import Chunk
from app.storage.material_storage import LocalMaterialStorage
from tests.conftest import CLEAN_DISTRACTORS, FakeLLM, good_question

LECTURE = """# Networks

Intro to the course.

## HTTP

HTTP is a request/response protocol. Status code 404 means the server cannot find the resource.

## DNS

DNS translates host names into IP addresses using a hierarchy of name servers.
"""


class FakeParser:
    def extract(self, path):
        return ParsedDocument(LECTURE, "fake")


FAKE_PARSERS = Parsers(pdf=FakeParser(), office=FakeParser())


def test_splitter_keeps_sections_apart_and_records_the_heading_path():
    chunks = split_markdown(LECTURE, "m1", min_chars=0)
    assert [c.source for c in chunks] == ["Networks", "Networks > HTTP", "Networks > DNS"]
    assert chunks[1].id == "m1:1"
    assert "404" in chunks[1].text and "404" not in chunks[2].text


def test_splitter_keeps_tiny_sections_selectable():
    chunks = split_markdown(LECTURE, "m1")
    assert [c.source for c in chunks] == ["Networks", "Networks > HTTP", "Networks > DNS"]


def test_splitter_merges_a_tiny_leftover_into_its_own_section():
    long_paragraph = "word " * 150
    chunks = split_markdown(f"# A\n\n{long_paragraph}\n\nshort tail", "m1", max_chars=800, min_chars=200)
    assert len(chunks) == 1 and chunks[0].text.endswith("short tail")


def test_splitter_cuts_long_sections_by_size_without_losing_text():
    body = "\n\n".join(f"Paragraph {i}. " + "word " * 100 for i in range(30))
    chunks = split_markdown(f"# Long\n\n{body}", "m1", max_chars=1000)
    assert len(chunks) > 3
    assert all(len(c.text) <= 1000 for c in chunks)
    assert all(f"Paragraph {i}." in "".join(c.text for c in chunks) for i in range(30))


def test_splitter_ignores_hash_lines_inside_code_fences():
    chunks = split_markdown("# Shell\n\n```\n# not a heading\nls\n```\n", "m1", min_chars=0)
    assert len(chunks) == 1 and "# not a heading" in chunks[0].text


def test_splitter_returns_nothing_for_empty_documents():
    assert split_markdown("  \n\n", "m1") == []


def test_retriever_ranks_the_relevant_chunk_first_and_drops_unrelated_ones():
    chunks = [
        Chunk(id="a", text="DNS translates host names into IP addresses"),
        Chunk(id="b", text="HTTP status code 404 means resource not found"),
        Chunk(id="c", text="Cooking pasta takes ten minutes"),
    ]
    found = LexicalRetriever().top_k("HTTP 404 status", chunks, k=3)
    assert [c.id for c in found] == ["b"]


def test_adapters_are_chosen_by_file_type():
    parsers = Parsers(pdf=PymupdfAdapter(), office=MarkItDownAdapter())
    assert isinstance(adapter_for("notes.MD", parsers), MarkdownAdapter)
    assert adapter_for("slides.pptx", parsers) is parsers.office
    assert adapter_for("paper.PDF", parsers) is parsers.pdf
    with pytest.raises(UnsupportedFormat):
        adapter_for("movie.mp4", parsers)


def test_storage_cannot_be_escaped_with_a_crafted_filename(tmp_path):
    storage = LocalMaterialStorage(tmp_path)
    saved = storage.save(7, "m1", "../../evil.md", b"x")
    assert saved == tmp_path / "7" / "m1" / "evil.md"
    assert not (tmp_path.parent / "evil.md").exists()


def make_client(tmp_path, llm, **settings_overrides):
    settings = Settings(
        provider="ollama", model="m", ollama_base_url="http://x", max_retries=1, materials_dir=tmp_path, **settings_overrides
    )
    return TestClient(create_app(settings, provider_factory=lambda config: llm, parsers=FAKE_PARSERS))


def upload(client, course_id=1, name="lecture.pdf"):
    return client.post("/materials", data={"course_id": course_id}, files={"file": (name, b"%PDF fake")})


def wait_for(client, material_id):
    for _ in range(50):
        material = client.get(f"/materials/{material_id}").json()
        if material["status"] != "PROCESSING":
            return material
        time.sleep(0.05)
    raise AssertionError("material never finished processing")


def test_uploaded_material_is_parsed_chunked_and_stored(tmp_path):
    client = make_client(tmp_path, FakeLLM([]))
    accepted = upload(client)
    assert accepted.status_code == 202

    material = wait_for(client, accepted.json()["id"])
    assert material["status"] == "READY" and material["chunk_count"] >= 1
    assert material["parser"] == "fake" and material["parse_seconds"] is not None
    folder = tmp_path / "1" / material["id"]
    assert (folder / "lecture.pdf").read_bytes() == b"%PDF fake"
    # the text lives in the database, not as loose files next to the original
    assert not (folder / "parsed.md").exists()
    assert client.get("/courses/1/materials").json()[0]["id"] == material["id"]
    assert client.get("/courses/2/materials").json() == []


def test_unsupported_upload_is_refused(tmp_path):
    assert upload(make_client(tmp_path, FakeLLM([])), name="movie.mp4").status_code == 415


def chunk_ids(client, material_id, *headings):
    """Ids of the material's pieces under the given headings, in reading order."""
    return [chunk["id"] for chunk in client.get(f"/materials/{material_id}/chunks").json() if chunk["heading"] in headings]


def test_generation_reads_the_pieces_it_is_given(tmp_path):
    llm = FakeLLM([good_question(), {"correct_supported": True, "correct_chunk_ids": [], "distractors_anchored": 3}, CLEAN_DISTRACTORS])
    client = make_client(tmp_path, llm)
    material_id = wait_for(client, upload(client).json()["id"])["id"]
    http = chunk_ids(client, material_id, "Networks > HTTP")
    llm.replies[1]["correct_chunk_ids"] = http

    job_id = client.post("/generate", json={"course_id": 1, "topic": "HTTP", "chunk_ids": http}).json()["id"]
    job = client.get(f"/jobs/{job_id}").json()

    assert job["status"] == "DONE"
    assert job["outcomes"][0]["status"] == "OK"
    assert job["outcomes"][0]["source_chunk_ids"] == http
    assert "404 means the server cannot find the resource" in llm.prompts[0]
    assert "DNS translates" not in llm.prompts[0] and "Intro to the course" not in llm.prompts[0]


def test_generation_rejects_pieces_of_another_course(tmp_path):
    client = make_client(tmp_path, FakeLLM([]))
    material_id = wait_for(client, upload(client, course_id=1).json()["id"])["id"]
    response = client.post("/generate", json={"course_id": 2, "topic": "HTTP", "chunk_ids": chunk_ids(client, material_id, "Networks > HTTP")})
    assert response.status_code == 422


def test_generation_rejects_pieces_that_no_longer_exist(tmp_path):
    client = make_client(tmp_path, FakeLLM([]))
    wait_for(client, upload(client).json()["id"])
    response = client.post("/generate", json={"course_id": 1, "topic": "HTTP", "chunk_ids": ["gone:7"]})
    assert response.status_code == 422
    assert "no longer exist" in response.json()["detail"]


def test_generation_needs_some_context(tmp_path):
    client = make_client(tmp_path, FakeLLM([]))
    assert client.post("/generate", json={"course_id": 1, "topic": "HTTP"}).status_code == 422


def test_original_is_deleted_when_not_kept_and_the_stored_text_is_reprocessed(tmp_path):
    client = make_client(tmp_path, FakeLLM([]), keep_originals=False)
    material = wait_for(client, upload(client).json()["id"])
    assert material["status"] == "READY"
    assert not (tmp_path / "1" / material["id"] / "lecture.pdf").exists()

    assert client.post(f"/materials/{material['id']}/reprocess").status_code == 202
    again = wait_for(client, material["id"])
    assert again["status"] == "READY" and again["chunk_count"] == material["chunk_count"]


def test_a_material_from_before_the_database_text_is_reprocessed_from_its_file(tmp_path):
    client = make_client(tmp_path, FakeLLM([]), keep_originals=False)
    material = wait_for(client, upload(client).json()["id"])
    legacy = tmp_path / "1" / material["id"] / "parsed.md"
    legacy.parent.mkdir(parents=True, exist_ok=True)
    legacy.write_text("# Old\n\n## Only section\n\n" + "text " * 60, encoding="utf-8")
    # drop the text the upload stored, as a material from the old version would be
    with client.app.state.materials._session.begin() as session:
        session.get(MaterialRow, material["id"]).markdown = None

    client.post(f"/materials/{material['id']}/reprocess")
    wait_for(client, material["id"])
    assert [c["heading"] for c in client.get(f"/materials/{material['id']}/chunks").json()] == ["Old > Only section"]


def test_the_pieces_of_a_material_are_listed_in_reading_order_with_their_headings(tmp_path):
    client = make_client(tmp_path, FakeLLM([]))
    material_id = wait_for(client, upload(client).json()["id"])["id"]
    chunks = client.get(f"/materials/{material_id}/chunks").json()
    assert [(c["position"], c["heading"]) for c in chunks] == [
        (0, "Networks"), (1, "Networks > HTTP"), (2, "Networks > DNS")
    ]
    assert "404 means" in chunks[1]["text"]
    assert client.get("/materials/nope/chunks").status_code == 404


GOOD_GROUNDING = {"correct_supported": True, "correct_chunk_ids": [], "distractors_anchored": 3}


def other_question(stem):
    question = good_question()
    question["stem"] = stem
    return question


def test_a_topic_with_more_text_than_fits_spreads_it_over_the_questions(tmp_path):
    llm = FakeLLM([
        other_question("Which protocol answers with status codes?"), GOOD_GROUNDING, CLEAN_DISTRACTORS,
        other_question("What do name servers translate?"), GOOD_GROUNDING, CLEAN_DISTRACTORS,
    ])
    # Room for one piece per prompt
    client = make_client(tmp_path, llm, context_chars=60)
    material_id = wait_for(client, upload(client).json()["id"])["id"]

    client.post("/generate", json={
        "course_id": 1, "topic": "Networks", "count": 2,
        "chunk_ids": chunk_ids(client, material_id, "Networks > HTTP", "Networks > DNS"),
    })

    first, second = llm.prompts[0], llm.prompts[3]
    assert "404" in first and "DNS translates" not in first
    assert "DNS translates" in second and "404" not in second


def test_a_topic_that_fits_gives_every_question_all_its_text(tmp_path):
    llm = FakeLLM([good_question(), GOOD_GROUNDING, CLEAN_DISTRACTORS])
    client = make_client(tmp_path, llm)
    material_id = wait_for(client, upload(client).json()["id"])["id"]

    client.post("/generate", json={
        "course_id": 1, "topic": "Networks", "chunk_ids": chunk_ids(client, material_id, "Networks > HTTP", "Networks > DNS"),
    })

    assert "404" in llm.prompts[0] and "DNS translates" in llm.prompts[0]


def test_the_focus_picks_the_most_relevant_pieces_first(tmp_path):
    llm = FakeLLM([good_question(), GOOD_GROUNDING, CLEAN_DISTRACTORS])
    client = make_client(tmp_path, llm, context_chars=60)
    material_id = wait_for(client, upload(client).json()["id"])["id"]

    client.post("/generate", json={
        "course_id": 1, "topic": "Networks", "focus": "name servers hierarchy",
        "chunk_ids": chunk_ids(client, material_id, "Networks > HTTP", "Networks > DNS"),
    })

    assert "DNS translates" in llm.prompts[0] and "404" not in llm.prompts[0]
    assert "Focus inside the topic: name servers hierarchy" in llm.prompts[0]


def test_a_parser_failure_is_reported_on_the_material(tmp_path):
    class Broken:
        def extract(self, path):
            raise RuntimeError("corrupt file")

    settings = Settings(provider="ollama", model="m", ollama_base_url="http://x", max_retries=1, materials_dir=tmp_path)
    client = TestClient(create_app(settings, parsers=Parsers(pdf=Broken(), office=Broken())))
    material = wait_for(client, upload(client).json()["id"])
    assert material["status"] == "FAILED" and "corrupt file" in material["error"]


def test_fallback_parser_takes_over_only_when_the_fast_one_finds_no_text(tmp_path):
    class Fast:
        def __init__(self, text):
            self.text = text

        def extract(self, path):
            return ParsedDocument(self.text, "fast")

    class Heavy:
        def extract(self, path):
            return ParsedDocument("recovered by ocr " * 10, "heavy")

    assert FallbackAdapter(Fast("x" * 100), Heavy()).extract(tmp_path).parser == "fast"
    assert FallbackAdapter(Fast("   "), Heavy()).extract(tmp_path).parser == "heavy"


def test_markitdown_reads_a_real_pptx(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "HTTP status codes"
    slide.placeholders[1].text = "404 means the server cannot find the resource"
    path = tmp_path / "lecture.pptx"
    deck.save(path)

    document = MarkItDownAdapter().extract(path)
    assert document.parser == "markitdown"
    assert "HTTP status codes" in document.markdown and "404 means" in document.markdown


def test_pymupdf_reads_a_real_pdf_and_finds_nothing_in_a_blank_one(tmp_path):
    import pymupdf

    with_text = pymupdf.open()
    with_text.new_page().insert_text((72, 72), "DNS translates host names into IP addresses")
    with_text.save(tmp_path / "text.pdf")
    blank = pymupdf.open()
    blank.new_page()
    blank.save(tmp_path / "blank.pdf")

    assert "DNS translates host names" in PymupdfAdapter().extract(tmp_path / "text.pdf").markdown
    assert PymupdfAdapter().extract(tmp_path / "blank.pdf").markdown.strip() == ""


def test_materials_and_jobs_survive_a_service_restart(tmp_path):
    database = f"sqlite:///{tmp_path / 'aqg.db'}"

    def start(llm):
        settings = Settings(
            provider="ollama", model="m", ollama_base_url="http://x", max_retries=1,
            materials_dir=tmp_path, database_url=database,
        )
        return TestClient(create_app(settings, provider_factory=lambda config: llm, parsers=FAKE_PARSERS))

    first = start(FakeLLM([good_question(), {"correct_supported": True, "correct_chunk_ids": [], "distractors_anchored": 3}, CLEAN_DISTRACTORS]))
    material_id = upload(first).json()["id"]
    wait_for(first, material_id)
    job_id = first.post(
        "/generate", json={"course_id": 1, "topic": "HTTP status 404", "chunks": [{"id": "c1", "text": "404 means not found"}]}
    ).json()["id"]

    second = start(FakeLLM([]))
    assert second.get(f"/materials/{material_id}").json()["status"] == "READY"
    assert second.get(f"/jobs/{job_id}").json()["status"] in ("DONE", "FAILED")
    # the chunks came back too: generation from the stored material still works
    response = second.post("/generate", json={"course_id": 1, "topic": "HTTP status 404", "chunk_ids": [f"{material_id}:1"]})
    assert response.status_code == 202
